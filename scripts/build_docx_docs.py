"""
SkillAlign Enterprise Markdown-to-DOCX Documentation Compiler (v2 — Consolidated)
Converts all Markdown documentation files in docs/ into 5 optimized, executive-grade
Microsoft Word (.docx) documents with rich typography, custom-styled tables, callout
blocks, code styling, and high-resolution visual diagrams embedded.

Consolidation Map:
  Doc 01: Vol 01 + Vol 08 -> Executive Overview & Architecture
  Doc 02: Vol 02 + Vol 03 -> Workflows & Functional Specification
  Doc 03: Vol 04 + Vol 05 + Vol 06 -> Matching Engine, Database & APIs
  Doc 04: Vol 07 -> Security & RBAC
  Doc 05: Vol 09 -> Demo Script & QA
"""

import os
import re
import glob
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

DOCS_DIR = os.path.join(os.path.dirname(__file__), "..", "docs")
DIAGRAMS_DIR = os.path.join(DOCS_DIR, "diagrams")

# --- Professional Enterprise Color Constants ---
COLOR_NAVY_DARK = RGBColor(15, 23, 42)
COLOR_NAVY_BLUE = RGBColor(30, 58, 138)
COLOR_ELECTRIC_BLUE = RGBColor(37, 99, 235)
COLOR_SLATE_DARK = RGBColor(51, 65, 85)
COLOR_SLATE_MED = RGBColor(100, 116, 139)
COLOR_TEXT_DARK = RGBColor(30, 41, 59)
COLOR_CODE_DARK = RGBColor(15, 23, 42)

HEX_HEADER_BG = "1E3A8A"
HEX_ZEBRA_BG = "F8FAFC"
HEX_WHITE_BG = "FFFFFF"
HEX_BORDER = "CBD5E1"
HEX_CALLOUT_BG = "EFF6FF"
HEX_CALLOUT_BORDER = "2563EB"
HEX_CODE_BG = "F1F5F9"
HEX_CODE_BORDER = "E2E8F0"

# --- Upgraded Font Sizes ---
FONT_TITLE = 28
FONT_H1 = 20
FONT_H2 = 17
FONT_H3 = 14
FONT_H4 = 12
FONT_BODY = 12
FONT_TABLE_HDR = 11
FONT_TABLE_BODY = 10.5
FONT_CODE = 10
FONT_CALLOUT = 11
FONT_CAPTION = 10
FONT_HEADER_FOOTER = 9
FONT_LINE_SPACING = 17


# =====================================================================
#  Low-Level Styling Helpers
# =====================================================================

def set_cell_shading(cell, color_hex):
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    cell._tc.get_or_add_tcPr().append(shd)

def set_cell_margins(cell, top=120, bottom=120, left=180, right=180):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'''
        <w:tcMar {nsdecls("w")}>
            <w:top w:w="{top}" w:type="dxa"/>
            <w:bottom w:w="{bottom}" w:type="dxa"/>
            <w:left w:w="{left}" w:type="dxa"/>
            <w:right w:w="{right}" w:type="dxa"/>
        </w:tcMar>
    ''')
    tcPr.append(tcMar)

def set_table_borders(table, color_hex=HEX_BORDER):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'''
        <w:tblBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="6" w:space="0" w:color="{color_hex}"/>
            <w:left w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:bottom w:val="single" w:sz="8" w:space="0" w:color="{color_hex}"/>
            <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:insideH w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>
            <w:insideV w:val="none" w:sz="0" w:space="0" w:color="auto"/>
        </w:tblBorders>
    ''')
    tblPr.append(borders)

def set_callout_borders(cell, color_hex=HEX_CALLOUT_BORDER, sz="24"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:left w:val="single" w:sz="{sz}" w:space="0" w:color="{color_hex}"/>
            <w:bottom w:val="none" w:sz="0" w:space="0" w:color="auto"/>
            <w:right w:val="none" w:sz="0" w:space="0" w:color="auto"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)

def set_code_borders(cell, color_hex=HEX_CODE_BORDER):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'''
        <w:tcBorders {nsdecls("w")}>
            <w:top w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>
            <w:left w:val="single" w:sz="8" w:space="0" w:color="{color_hex}"/>
            <w:bottom w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>
            <w:right w:val="single" w:sz="4" w:space="0" w:color="{color_hex}"/>
        </w:tcBorders>
    ''')
    tcPr.append(borders)


# =====================================================================
#  Inline Formatting
# =====================================================================

def format_inline_text(paragraph, text, default_font="Calibri", default_size=FONT_BODY, default_color=COLOR_TEXT_DARK):
    tokens = re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\[[^\]]+\]\([^)]+\)|\$[^\$]+\$)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            inner = token[2:-2]
            run = paragraph.add_run(inner)
            run.bold = True
            run.font.name = default_font
            run.font.size = Pt(default_size)
            run.font.color.rgb = default_color
        elif token.startswith('*') and token.endswith('*'):
            inner = token[1:-1]
            run = paragraph.add_run(inner)
            run.italic = True
            run.font.name = default_font
            run.font.size = Pt(default_size)
            run.font.color.rgb = default_color
        elif token.startswith('`') and token.endswith('`'):
            inner = token[1:-1]
            run = paragraph.add_run(inner)
            run.font.name = "Consolas"
            run.font.size = Pt(default_size * 0.85)
            run.font.color.rgb = COLOR_NAVY_BLUE
            run.bold = True
        elif token.startswith('$') and token.endswith('$'):
            inner = token[1:-1]
            run = paragraph.add_run(inner)
            run.font.name = "Cambria Math"
            run.font.size = Pt(default_size)
            run.font.color.rgb = COLOR_NAVY_DARK
            run.italic = True
        elif token.startswith('[') and ']' in token and '(' in token and token.endswith(')'):
            m = re.match(r'\[([^\]]+)\]\(([^)]+)\)', token)
            if m:
                link_text = m.group(1)
                run = paragraph.add_run(link_text)
                run.font.name = default_font
                run.font.size = Pt(default_size)
                run.font.color.rgb = COLOR_ELECTRIC_BLUE
                run.underline = True
            else:
                run = paragraph.add_run(token)
                run.font.name = default_font
                run.font.size = Pt(default_size)
        else:
            run = paragraph.add_run(token)
            run.font.name = default_font
            run.font.size = Pt(default_size)
            run.font.color.rgb = default_color


# =====================================================================
#  Document Structure Elements
# =====================================================================

def add_header_footer(doc, title_text):
    for section in doc.sections:
        section.top_margin = Inches(0.9)
        section.bottom_margin = Inches(0.9)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)

        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run(f"SkillAlign Platform  |  {title_text}")
        hrun.font.name = "Calibri"
        hrun.font.size = Pt(FONT_HEADER_FOOTER)
        hrun.font.color.rgb = COLOR_SLATE_MED

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("SkillAlign Enterprise Engineering Documentation  -  Confidential & Proprietary")
        frun.font.name = "Calibri"
        frun.font.size = Pt(FONT_HEADER_FOOTER)
        frun.font.color.rgb = COLOR_SLATE_MED


def create_document_cover(doc, title, subtitle=None, metadata_dict=None):
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_title.paragraph_format.space_before = Pt(18)
    p_title.paragraph_format.space_after = Pt(6)
    run_title = p_title.add_run(title)
    run_title.font.name = "Calibri"
    run_title.font.size = Pt(FONT_TITLE)
    run_title.bold = True
    run_title.font.color.rgb = COLOR_NAVY_BLUE

    if subtitle:
        p_sub = doc.add_paragraph()
        p_sub.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p_sub.paragraph_format.space_after = Pt(14)
        run_sub = p_sub.add_run(subtitle)
        run_sub.font.name = "Calibri"
        run_sub.font.size = Pt(13)
        run_sub.italic = True
        run_sub.font.color.rgb = COLOR_SLATE_MED

    if metadata_dict:
        table = doc.add_table(rows=1, cols=len(metadata_dict))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, "E2E8F0")
        col_widths = [Inches(6.5 / len(metadata_dict))] * len(metadata_dict)

        for idx, (k, v) in enumerate(metadata_dict.items()):
            cell = table.rows[0].cells[idx]
            cell.width = col_widths[idx]
            set_cell_shading(cell, "F8FAFC")
            set_cell_margins(cell, top=100, bottom=100, left=120, right=120)

            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_after = Pt(0)

            r_k = p.add_run(f"{k.upper()}\n")
            r_k.font.name = "Calibri"
            r_k.font.size = Pt(8.5)
            r_k.bold = True
            r_k.font.color.rgb = COLOR_SLATE_MED

            r_v = p.add_run(v)
            r_v.font.name = "Calibri"
            r_v.font.size = Pt(11)
            r_v.bold = True
            r_v.font.color.rgb = COLOR_NAVY_DARK

        p_space = doc.add_paragraph()
        p_space.paragraph_format.space_before = Pt(8)
        p_space.paragraph_format.space_after = Pt(8)


# =====================================================================
#  Content Elements: Tables, Callouts, Code, Diagrams
# =====================================================================

def insert_styled_table(doc, headers, rows):
    if not headers or not rows:
        return
    table = doc.add_table(rows=len(rows) + 1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders(table, HEX_BORDER)

    total_width = Inches(6.5)
    col_width = total_width / len(headers)

    # Header Row
    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
        hdr_cells[i].width = col_width
        set_cell_shading(hdr_cells[i], HEX_HEADER_BG)
        set_cell_margins(hdr_cells[i], top=140, bottom=140, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(header_text)
        run.font.name = "Calibri"
        run.font.size = Pt(FONT_TABLE_HDR)
        run.bold = True
        run.font.color.rgb = RGBColor(255, 255, 255)

    # Data Rows
    for r_idx, row_data in enumerate(rows):
        row_cells = table.rows[r_idx + 1].cells
        bg_hex = HEX_ZEBRA_BG if r_idx % 2 == 1 else HEX_WHITE_BG
        for c_idx, cell_value in enumerate(row_data):
            if c_idx < len(row_cells):
                cell = row_cells[c_idx]
                cell.width = col_width
                set_cell_shading(cell, bg_hex)
                set_cell_margins(cell, top=110, bottom=110, left=140, right=140)
                p = cell.paragraphs[0]
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.space_after = Pt(0)
                format_inline_text(p, cell_value, default_size=FONT_TABLE_BODY)

    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_after = Pt(6)


def insert_callout_box(doc, text, alert_type="NOTE"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)

    bg_color = HEX_CALLOUT_BG
    border_color = HEX_CALLOUT_BORDER
    title_color = COLOR_ELECTRIC_BLUE
    icon = "NOTE"

    if "IMPORTANT" in alert_type:
        bg_color = "FEF3C7"
        border_color = "D97706"
        title_color = RGBColor(217, 119, 6)
        icon = "IMPORTANT"
    elif "TIP" in alert_type:
        bg_color = "F0FDF4"
        border_color = "059669"
        title_color = RGBColor(5, 150, 105)
        icon = "TIP"
    elif "WARNING" in alert_type or "CAUTION" in alert_type:
        bg_color = "FFF1F2"
        border_color = "E11D48"
        title_color = RGBColor(225, 29, 72)
        icon = "WARNING"

    set_cell_shading(cell, bg_color)
    set_callout_borders(cell, border_color, sz="24")
    set_cell_margins(cell, top=120, bottom=120, left=180, right=160)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(2)
    r_hdr = p.add_run(f"{icon}: ")
    r_hdr.font.name = "Calibri"
    r_hdr.font.size = Pt(FONT_CALLOUT)
    r_hdr.bold = True
    r_hdr.font.color.rgb = title_color

    format_inline_text(p, text, default_size=FONT_CALLOUT, default_color=COLOR_NAVY_DARK)

    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_after = Pt(6)


def insert_code_block(doc, code_text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.rows[0].cells[0]
    cell.width = Inches(6.5)

    set_cell_shading(cell, HEX_CODE_BG)
    set_code_borders(cell, HEX_CODE_BORDER)
    set_cell_margins(cell, top=120, bottom=120, left=160, right=160)

    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(13)

    run = p.add_run(code_text.strip())
    run.font.name = "Consolas"
    run.font.size = Pt(FONT_CODE)
    run.font.color.rgb = COLOR_CODE_DARK

    p_post = doc.add_paragraph()
    p_post.paragraph_format.space_after = Pt(6)


def insert_diagram(doc, img_filename, caption_text):
    img_path = os.path.join(DIAGRAMS_DIR, img_filename)
    if os.path.exists(img_path):
        p_img = doc.add_paragraph()
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_img.paragraph_format.space_before = Pt(10)
        p_img.paragraph_format.space_after = Pt(4)
        p_img.add_run().add_picture(img_path, width=Inches(6.3))

        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_cap.paragraph_format.space_after = Pt(12)
        r_cap = p_cap.add_run(f"Figure: {caption_text}")
        r_cap.font.name = "Calibri"
        r_cap.font.size = Pt(FONT_CAPTION)
        r_cap.font.color.rgb = COLOR_SLATE_MED
        r_cap.italic = True
    else:
        print(f"  [WARN] Diagram image not found: {img_path}")


# =====================================================================
#  Markdown Parser -> DOCX
# =====================================================================

def parse_markdown_content_into_doc(doc, md_content, diagram_mappings=None, heading_diagram_hooks=None):
    lines = md_content.splitlines()
    i = 0
    in_code_block = False
    code_lang = ""
    code_lines = []

    while i < len(lines):
        line = lines[i].rstrip()

        # --- Code / Mermaid Blocks ---
        if line.startswith("```"):
            if in_code_block:
                full_code = "\n".join(code_lines)
                matched_fig = None

                if code_lang == "mermaid" or any(c in full_code for c in ["flowchart", "sequenceDiagram", "erDiagram"]):
                    if diagram_mappings:
                        context_window = "\n".join(lines[max(0, i-30):i+5]).lower()
                        for key, (img, cap) in diagram_mappings.items():
                            if key.lower() in full_code.lower() or key.lower() in context_window:
                                matched_fig = (img, cap)
                                break
                    if matched_fig:
                        insert_diagram(doc, matched_fig[0], matched_fig[1])
                    else:
                        insert_code_block(doc, full_code)
                else:
                    insert_code_block(doc, full_code)
                in_code_block = False
                code_lines = []
                code_lang = ""
            else:
                in_code_block = True
                code_lang = line[3:].strip().lower()
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(lines[i].rstrip('\n'))
            i += 1
            continue

        # --- Tables ---
        if "|" in line and not line.startswith("```"):
            table_lines = []
            while i < len(lines) and "|" in lines[i].rstrip() and not lines[i].startswith("```"):
                table_lines.append(lines[i].rstrip())
                i += 1

            if len(table_lines) >= 2:
                raw_headers = [c.strip() for c in table_lines[0].split("|") if c.strip()]
                data_rows = []
                for row_line in table_lines[2:]:
                    cols = [c.strip() for c in row_line.split("|")]
                    if len(cols) >= 2 and cols[0] == "":
                        cols = cols[1:]
                    if len(cols) >= 1 and cols[-1] == "":
                        cols = cols[:-1]
                    cols = [c.strip() for c in cols]
                    if cols:
                        data_rows.append(cols)
                insert_styled_table(doc, raw_headers, data_rows)
            continue

        # --- Blockquotes / Alerts ---
        if line.startswith(">"):
            alert_match = re.match(r'>\s*\[!([A-Z]+)\]\s*(.*)', line)
            if alert_match:
                alert_type = alert_match.group(1)
                alert_text = alert_match.group(2)
                i += 1
                while i < len(lines) and lines[i].startswith(">"):
                    cont = lines[i].lstrip(">").strip()
                    if cont:
                        alert_text += " " + cont
                    i += 1
                insert_callout_box(doc, alert_text, alert_type)
                continue
            else:
                quote_text = line.lstrip(">").strip()
                i += 1
                while i < len(lines) and lines[i].startswith(">"):
                    quote_text += " " + lines[i].lstrip(">").strip()
                    i += 1
                insert_callout_box(doc, quote_text, "NOTE")
                continue

        # --- Headings ---
        if line.startswith("# "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line[2:].strip())
            run.font.name = "Calibri"
            run.font.size = Pt(FONT_H1)
            run.bold = True
            run.font.color.rgb = COLOR_NAVY_BLUE
        elif line.startswith("## "):
            h_text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(h_text)
            run.font.name = "Calibri"
            run.font.size = Pt(FONT_H2)
            run.bold = True
            run.font.color.rgb = COLOR_NAVY_BLUE

            if heading_diagram_hooks:
                for hook_key, (img, cap) in heading_diagram_hooks.items():
                    if hook_key.lower() in h_text.lower():
                        insert_diagram(doc, img, cap)
        elif line.startswith("### "):
            h_text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(12)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(h_text)
            run.font.name = "Calibri"
            run.font.size = Pt(FONT_H3)
            run.bold = True
            run.font.color.rgb = COLOR_ELECTRIC_BLUE

            if heading_diagram_hooks:
                for hook_key, (img, cap) in heading_diagram_hooks.items():
                    if hook_key.lower() in h_text.lower():
                        insert_diagram(doc, img, cap)
        elif line.startswith("#### "):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[5:].strip())
            run.font.name = "Calibri"
            run.font.size = Pt(FONT_H4)
            run.bold = True
            run.font.color.rgb = COLOR_SLATE_DARK
        elif line.startswith("---"):
            pass
        elif line.startswith("- ") or line.startswith("* "):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = Pt(FONT_LINE_SPACING)
            format_inline_text(p, line[2:].strip())
        elif re.match(r'^\d+\.\s', line):
            m = re.match(r'^(\d+\.)\s*(.*)', line)
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_after = Pt(3)
            p.paragraph_format.line_spacing = Pt(FONT_LINE_SPACING)
            format_inline_text(p, m.group(2).strip())
        elif line.strip() == "":
            pass
        else:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(5)
            p.paragraph_format.line_spacing = Pt(FONT_LINE_SPACING)
            format_inline_text(p, line)

        i += 1


# =====================================================================
#  Section Divider (between merged volumes)
# =====================================================================

def add_section_divider(doc, section_title):
    """Add a visual page break and section title between merged volumes."""
    doc.add_page_break()
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(24)
    p.paragraph_format.space_after = Pt(12)
    run = p.add_run(section_title)
    run.font.name = "Calibri"
    run.font.size = Pt(22)
    run.bold = True
    run.font.color.rgb = COLOR_NAVY_BLUE

    # Add a thin horizontal rule
    p_rule = doc.add_paragraph()
    p_rule.paragraph_format.space_after = Pt(8)
    # Use a bottom border on the paragraph as a visual rule
    pPr = p_rule._p.get_or_add_pPr()
    pBdr = parse_xml(f'''
        <w:pBdr {nsdecls("w")}>
            <w:bottom w:val="single" w:sz="6" w:space="1" w:color="{HEX_BORDER}"/>
        </w:pBdr>
    ''')
    pPr.append(pBdr)


# =====================================================================
#  Build a Consolidated Document from Multiple MD Sources
# =====================================================================

def build_consolidated_doc(doc_config):
    """Build a single DOCX from one or more markdown source files."""
    docx_path = os.path.join(DOCS_DIR, doc_config["docx"])
    print(f"\n  Building: {doc_config['docx']}")

    doc = Document()
    add_header_footer(doc, doc_config["title"])

    metadata = {
        "Document": doc_config["title"],
        "Author": "SkillAlign Architecture Team",
        "Status": "Production Verified",
        "Version": "2.0 Consolidated"
    }
    create_document_cover(doc, doc_config["title"], doc_config.get("subtitle", "Engineering Technical Specification"), metadata)

    sources = doc_config["sources"]
    for s_idx, source in enumerate(sources):
        md_path = os.path.join(DOCS_DIR, source["md"])
        if not os.path.exists(md_path):
            print(f"  [WARN] Source not found: {source['md']}")
            continue

        # Add section divider between merged volumes (skip for first)
        if s_idx > 0 and source.get("section_title"):
            add_section_divider(doc, source["section_title"])

        with open(md_path, "r", encoding="utf-8") as f:
            md_content = f.read()

        print(f"    + Parsing: {source['md']}")
        parse_markdown_content_into_doc(
            doc, md_content,
            source.get("diagrams"),
            source.get("heading_hooks")
        )

    doc.save(docx_path)
    print(f"  [OK] Saved: {doc_config['docx']}")


# =====================================================================
#  Consolidated Document Configurations (9 volumes -> 5 docs)
# =====================================================================

CONSOLIDATED_CONFIGS = [
    # ---------------------------------------------------------------
    # DOC 01: Executive Overview & Architecture (Vol 01 + Vol 08)
    # ---------------------------------------------------------------
    {
        "docx": "01_Executive_Overview_and_Architecture.docx",
        "title": "Executive Overview & System Architecture",
        "subtitle": "Platform Overview, Technology Stack, Architecture Layers, Codebase Structure & Deployment",
        "sources": [
            {
                "md": "01_EXECUTIVE_OVERVIEW_AND_SYSTEM_ARCHITECTURE.md",
                "diagrams": {
                    "Presentation Layer": ("fig_01_arch_layers.png", "SkillAlign 6-Tier Layered Architecture"),
                    "Technology Status": ("fig_00_platform_overview.png", "SkillAlign Enterprise Platform Ecosystem"),
                },
                "heading_hooks": {
                    "3. System Architecture Layers": ("fig_01_arch_layers.png", "SkillAlign 6-Tier Layered System Architecture")
                }
            },
            {
                "md": "08_FRONTEND_AND_BACKEND_CODEBASE_STRUCTURE.md",
                "section_title": "Codebase Structure, Deployment & Engineering Decisions",
                "diagrams": {
                    "Deployment Architecture": ("fig_08_deployment_topology.png", "Production Deployment & Infrastructure Topology"),
                    "subgraph": ("fig_08_deployment_topology.png", "Containerized Infrastructure & Cluster Topology"),
                    "flowchart TD": ("fig_08_deployment_topology.png", "Infrastructure & Cluster Topology")
                },
                "heading_hooks": {
                    "3. Infrastructure & Deployment Architecture": ("fig_08_deployment_topology.png", "Production Deployment & Infrastructure Topology")
                }
            }
        ]
    },
    # ---------------------------------------------------------------
    # DOC 02: Workflows & Functional Specification (Vol 02 + Vol 03)
    # ---------------------------------------------------------------
    {
        "docx": "02_Workflows_and_Functional_Specification.docx",
        "title": "Workflows, State Machines & Functional Specification",
        "subtitle": "End-to-End Recruitment Workflows, 21-State Machine & 26 Feature Specifications",
        "sources": [
            {
                "md": "02_END_TO_END_WORKFLOWS_AND_STATE_MACHINES.md",
                "diagrams": {
                    "Candidate Signs Up": ("fig_02_recruitment_lifecycle.png", "Complete End-to-End Recruitment Workflow"),
                    "Visit Landing Page": ("fig_02_role_workflows.png", "Role-Specific Recruitment Workflows"),
                    "Login as HR": ("fig_02_role_workflows.png", "Role-Specific Recruitment Workflows"),
                    "Login as Recruiter": ("fig_02_role_workflows.png", "Role-Specific Recruitment Workflows"),
                    "State Machine": ("fig_02_state_machine.png", "21-State Recruitment State Machine"),
                    "flowchart TD": ("fig_02_recruitment_lifecycle.png", "Recruitment Lifecycle Flowchart")
                },
                "heading_hooks": {
                    "3. Enterprise Recruitment State Machine": ("fig_02_state_machine.png", "21-State Recruitment Lifecycle State Machine")
                }
            },
            {
                "md": "03_FUNCTIONAL_SPECIFICATION_DOCUMENT.md",
                "section_title": "Functional Specification Document (FSD) & Use Cases",
                "diagrams": {
                    "Feature Inventory": ("fig_03_fsd_feature_map.png", "Functional Feature Inventory & Module Flow"),
                    "F-004": ("fig_04_resume_pipeline.png", "Resume Upload & Text Extraction Pipeline"),
                    "F-018": ("fig_07_seq_offer_generation.png", "Offer Letter Drafting Flow")
                },
                "heading_hooks": {
                    "1. Feature Inventory & Implementation Status Matrix": ("fig_03_fsd_feature_map.png", "Functional Feature Map & Architecture"),
                    "F-004: Resume Document Upload & Multi-Format Text Extraction": ("fig_04_resume_pipeline.png", "Resume Extraction Pipeline"),
                    "F-018: Compensation Discussion & Offer Drafting": ("fig_07_seq_offer_generation.png", "Offer Generation Flow")
                }
            }
        ]
    },
    # ---------------------------------------------------------------
    # DOC 03: Matching Engine, Database & APIs (Vol 04 + Vol 05 + Vol 06)
    # ---------------------------------------------------------------
    {
        "docx": "03_Matching_Engine_Database_and_APIs.docx",
        "title": "Matching Engine, Database Design & API Inventory",
        "subtitle": "Multi-Factor Scoring, Resume Pipeline, Database Schema, ERD & 54 REST API Endpoints",
        "sources": [
            {
                "md": "04_MATCHING_ENGINE_AND_RESUME_PROCESSING.md",
                "diagrams": {
                    "Dimension": ("fig_04_matching_engine.png", "Multi-Factor Scoring Dimensions"),
                    "SkillAlign Candidate Fit Score": ("fig_04_matching_engine.png", "Candidate Fit Scoring"),
                    "pdfplumber": ("fig_04_resume_pipeline.png", "Resume Extraction Pipeline")
                },
                "heading_hooks": {
                    "1. Multi-Factor Matching Engine": ("fig_04_matching_engine.png", "Multi-Factor Fit Scoring Weight Distribution"),
                    "2. Deterministic Resume Extraction & Parsing Pipeline": ("fig_04_resume_pipeline.png", "Resume Ingestion & Parsing Pipeline")
                }
            },
            {
                "md": "05_DATABASE_DESIGN_AND_ERD.md",
                "section_title": "Database Design, Schema & Entity-Relationship Diagrams",
                "diagrams": {
                    "erDiagram": ("fig_05_erd.png", "Complete Entity-Relationship Diagram (23 Models)"),
                    "ERD": ("fig_05_erd.png", "Database Schema & Entity Relationships")
                },
                "heading_hooks": {
                    "2. Comprehensive Entity-Relationship Diagram (ERD)": ("fig_05_erd.png", "Complete Entity-Relationship Diagram")
                }
            },
            {
                "md": "06_API_INVENTORY_AND_INTEGRATION_GUIDE.md",
                "section_title": "API Inventory, Integration & Traceability",
                "diagrams": {
                    "Master API Inventory": ("fig_06_api_gateway.png", "FastAPI REST API Routing Flow"),
                    "Axios": ("fig_06_api_gateway.png", "FastAPI REST API Architecture"),
                    "flowchart LR": ("fig_06_api_gateway.png", "Frontend-Backend API Integration Flow")
                },
                "heading_hooks": {
                    "1. Master API Inventory": ("fig_06_api_gateway.png", "FastAPI REST API Gateway & Routing Flow")
                }
            }
        ]
    },
    # ---------------------------------------------------------------
    # DOC 04: Security & RBAC (Vol 07 — standalone, dense content)
    # ---------------------------------------------------------------
    {
        "docx": "04_Security_and_RBAC.docx",
        "title": "Security Architecture, RBAC & Sequence Diagrams",
        "subtitle": "7-Layer Defense-in-Depth, Role-Based Access Control & Core Lifecycle Sequences",
        "sources": [
            {
                "md": "07_SECURITY_ERROR_HANDLING_AND_RBAC.md",
                "diagrams": {
                    "Layered Security": ("fig_07_security_layers.png", "7-Layer Defense-in-Depth Architecture"),
                    "User Authentication": ("fig_07_seq_auth_jwt.png", "User Authentication Sequence Flow"),
                    "Resume Upload": ("fig_04_resume_pipeline.png", "Resume Ingestion Pipeline"),
                    "Offer Letter": ("fig_07_seq_offer_generation.png", "Offer Generation Sequence"),
                    "sequenceDiagram": ("fig_07_seq_offer_generation.png", "Core Lifecycle Sequence")
                },
                "heading_hooks": {
                    "2. Security Architecture Deep Dive": ("fig_07_security_layers.png", "Multi-Layered Security Architecture"),
                    "4.1 Sequence Diagram: User Authentication & JWT Issuance": ("fig_07_seq_auth_jwt.png", "User Authentication & JWT Sequence"),
                    "4.7 Sequence Diagram: Formal Offer Letter Generation & Approval": ("fig_07_seq_offer_generation.png", "Offer Letter Generation Sequence")
                }
            }
        ]
    },
    # ---------------------------------------------------------------
    # DOC 05: Demo Script & QA (Vol 09)
    # ---------------------------------------------------------------
    {
        "docx": "05_Demo_Script_and_QA.docx",
        "title": "Technical Demo Script, Postman Flow & Engineering Q&A",
        "subtitle": "15-Step Live API Demo, Timed Presentation Script & 25+ Technical Defense Q&As",
        "sources": [
            {
                "md": "09_ENGINEERING_DEMO_SCRIPT_AND_QA.md",
                "diagrams": {
                    "flowchart LR": ("fig_09_api_demo_flow.png", "15-Step Live API Demo Verification Sequence"),
                    "Postman": ("fig_09_api_demo_flow.png", "Step-by-Step API Demo Flow")
                },
                "heading_hooks": {
                    "1. Step-by-Step Postman / API Demo Flow": ("fig_09_api_demo_flow.png", "15-Step Live API & Postman Demo Sequence")
                }
            }
        ]
    }
]


# =====================================================================
#  Cleanup Old Files
# =====================================================================

OLD_DOCX_FILES = [
    "01_EXECUTIVE_OVERVIEW_AND_SYSTEM_ARCHITECTURE.docx",
    "02_END_TO_END_WORKFLOWS_AND_STATE_MACHINES.docx",
    "03_FUNCTIONAL_SPECIFICATION_DOCUMENT.docx",
    "04_MATCHING_ENGINE_AND_RESUME_PROCESSING.docx",
    "05_DATABASE_DESIGN_AND_ERD.docx",
    "06_API_INVENTORY_AND_INTEGRATION_GUIDE.docx",
    "07_SECURITY_ERROR_HANDLING_AND_RBAC.docx",
    "08_FRONTEND_AND_BACKEND_CODEBASE_STRUCTURE.docx",
    "09_ENGINEERING_DEMO_SCRIPT_AND_QA.docx",
    "SkillAlign_Complete_System_Specification.docx",
    "SkillAlign_Master_Documentation_Guide.docx",
]

def cleanup_old_docx():
    """Remove old individual DOCX files that have been consolidated."""
    print("\n  Cleaning up old DOCX files...")
    removed = 0
    for old_file in OLD_DOCX_FILES:
        path = os.path.join(DOCS_DIR, old_file)
        if os.path.exists(path):
            os.remove(path)
            print(f"    [-] Removed: {old_file}")
            removed += 1
    print(f"  [OK] Cleaned up {removed} old files.")


# =====================================================================
#  Main Entry Point
# =====================================================================

def main():
    print("=" * 62)
    print("  SkillAlign DOCX Compiler v2 (Consolidated Edition)")
    print("=" * 62)
    print(f"  Source: {os.path.abspath(DOCS_DIR)}")
    print(f"  Diagrams: {os.path.abspath(DIAGRAMS_DIR)}")
    print(f"  Output: 5 consolidated DOCX documents")
    print("-" * 62)

    # Step 1: Remove old files
    cleanup_old_docx()

    # Step 2: Build 5 consolidated documents
    print("\n  Building 5 consolidated documents...\n")
    for cfg in CONSOLIDATED_CONFIGS:
        build_consolidated_doc(cfg)

    print("\n" + "=" * 62)
    print("  [SUCCESS] All 5 consolidated DOCX documents built!")
    print("=" * 62)
    print("\n  Output files:")
    for cfg in CONSOLIDATED_CONFIGS:
        docx_path = os.path.join(DOCS_DIR, cfg["docx"])
        if os.path.exists(docx_path):
            size_kb = os.path.getsize(docx_path) / 1024
            print(f"    {cfg['docx']}  ({size_kb:.0f} KB)")


if __name__ == "__main__":
    main()
